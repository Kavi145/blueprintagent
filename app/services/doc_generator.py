import io
from docx import Document

def create_blueprint_docx(plan: dict) -> io.BytesIO:
    doc = Document()
    doc.add_heading(plan.get("title", "Project Blueprint"), 0)

    doc.add_heading("1. Objectives", level=1)
    for obj in plan.get("objectives", []):
        doc.add_paragraph(obj, style='List Bullet')

    doc.add_heading("2. Requirements", level=1)
    for req in plan.get("requirements", []):
        doc.add_paragraph(req, style='List Bullet')

    doc.add_heading("3. Execution Tasks & Dependencies", level=1)
    for task in plan.get("tasks", []):
        p = doc.add_paragraph(style='List Number')
        p.add_run(f"**{task.get('name')}** (Timeline: {task.get('timeline')})\n").bold = True
        p.add_run(f"Description: {task.get('description')}\n")
        p.add_run(f"Dependencies: {', '.join(map(str, task.get('dependencies', [])) or ['None'])}")

    doc.add_heading("4. Resources", level=1)
    for res in plan.get("resources", []):
        doc.add_paragraph(res, style='List Bullet')

    doc.add_heading("5. Risks & Mitigations", level=1)
    for risk in plan.get("risks", []):
        doc.add_paragraph(risk, style='List Bullet')

    doc.add_heading("6. Execution Strategy", level=1)
    doc.add_paragraph(plan.get("execution_strategy", plan.get("execution_schema", "")))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer